"""
Document Extraction & Processing Pipeline
Target: Freesbe Document Intelligence Task (Vehicle Financing & Registrations)
Dependencies: pydantic, python-dotenv, google-genai (for live API calls)
"""

import os
import sys
import json
from datetime import date
from enum import Enum
from typing import Any, Dict, List, Optional
from pydantic import BaseModel, Field, ValidationError, field_validator, model_validator
from google import genai

types = genai.types


# Ensure UTF-8 output on Windows terminals
reconfigure_stdout = getattr(sys.stdout, "reconfigure", None)
if callable(reconfigure_stdout):
    reconfigure_stdout(encoding="utf-8")

# Automatically load environment variables from .env if available
try:
    import dotenv
    dotenv.load_dotenv()
except ImportError:
    pass


# ==========================================
# 1. Domain Entities & Schemas
# ==========================================
class DocumentCategory(str, Enum):
    VEHICLE_LICENSE = "Vehicle License"
    FINANCING_CONTRACT = "Financing Contract"
    DAMAGE_REPORT = "Damage Report"
    UNKNOWN = "Unknown"


class ProcessingStatus(str, Enum):
    APPROVED = "APPROVED"
    NEEDS_REVIEW = "NEEDS_REVIEW"
    REJECTED = "REJECTED"


class ExtractedVehicleDocument(BaseModel):
    """
    Data contract enforcing structural integrity across both
    LLM Structured Outputs and downstream business systems.
    """
    document_id: str = Field(..., description="Unique document tracking identifier")
    license_plate: str = Field(..., description="Israeli vehicle registration plate")
    customer_name: str = Field(..., description="Full legal name of customer or owner")
    customer_id: Optional[str] = Field(None, description="Israeli National ID (Teudat Zehut) - 9 digits")
    deal_amount: Optional[float] = Field(None, description="Transaction or financing amount in ILS")
    contract_date: Optional[date] = Field(None, description="Document issue or signature date")
    category: DocumentCategory = DocumentCategory.UNKNOWN
    estimated_damage: Optional[float] = Field(None, description="Estimated damage cost for Damage Reports")

    @field_validator("license_plate")
    @classmethod
    def clean_and_validate_plate(cls, value: str) -> str:
        """
        Normalizes Israeli license plates by removing dashes and whitespace,
        verifying the plate strictly comprises 7 or 8 numeric digits.
        """
        cleaned = value.replace("-", "").replace(" ", "").strip()
        if not cleaned.isdigit() or len(cleaned) not in (7, 8):
            raise ValueError(
                f"Invalid Israeli license plate format: '{value}'. Expected 7 or 8 numeric digits."
            )
        return cleaned

    @field_validator("customer_id")
    @classmethod
    def validate_israeli_id(cls, value: Optional[str]) -> Optional[str]:
        """
        Validates Israeli National ID (Teudat Zehut) using the official modulo-10 algorithm.
        Accepts 8 or 9 digits and zero-pads to 9 digits before verifying.
        """
        if value is None:
            return None
        cleaned = value.replace("-", "").replace(" ", "").strip().zfill(9)
        if not cleaned.isdigit() or len(cleaned) != 9:
            raise ValueError(f"Israeli ID must contain exactly 9 numeric digits: '{value}'")
        total = sum(
            (val - 9 if val > 9 else val)
            for i, digit in enumerate(cleaned)
            for val in [int(digit) * ((i % 2) + 1)]
        )
        if total % 10 != 0:
            raise ValueError(f"Invalid Israeli ID checksum: '{value}'")
        return cleaned

    @field_validator("deal_amount", "estimated_damage")
    @classmethod
    def validate_positive_amount(cls, value: Optional[float]) -> Optional[float]:
        """Ensures monetary values are strictly positive."""
        if value is not None and value <= 0:
            raise ValueError(f"Monetary amount must be greater than zero. Received: {value}")
        return value

    @model_validator(mode="after")
    def verify_cross_field_rules(self):
        """Cross-field business rule: Damage Reports must specify estimated_damage."""
        if self.category == DocumentCategory.DAMAGE_REPORT and self.estimated_damage is None:
            raise ValueError("Damage Reports require an 'estimated_damage' value.")
        return self


# ==========================================
# 2. Validation & Routing Engine
# ==========================================
def process_extracted_document(raw_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Core post-processing function:
    1. Fallback classification via content heuristics if category was not determined.
    2. Data contract enforcement and normalization via Pydantic.
    3. Missing optional field detection.
    4. Operational triage into APPROVED, NEEDS_REVIEW, or REJECTED.
    """
    snippet = raw_data.get("raw_text_snippet", "").lower()
    inferred_category = DocumentCategory.UNKNOWN

    if any(k in snippet for k in ["רכב", "רישיון", "שנת ייצור", "תוצר"]):
        inferred_category = DocumentCategory.VEHICLE_LICENSE
    elif any(k in snippet for k in ["מימון", "הלוואה", "לוח סילוקין", "תשלומים"]):
        inferred_category = DocumentCategory.FINANCING_CONTRACT
    elif any(k in snippet for k in ["שמאי", "נזק", "אומדן", "דו\"ח שמאי"]):
        inferred_category = DocumentCategory.DAMAGE_REPORT

    payload = dict(raw_data)
    if not payload.get("category") or payload.get("category") == DocumentCategory.UNKNOWN.value:
        payload["category"] = inferred_category.value

    try:
        doc = ExtractedVehicleDocument(**payload)

        # Identify missing non-critical optional attributes
        missing_optional: List[str] = []
        if doc.customer_id is None:
            missing_optional.append("customer_id")
        if doc.deal_amount is None:
            missing_optional.append("deal_amount")
        if doc.contract_date is None:
            missing_optional.append("contract_date")

        # Business triage: missing fields or high-value contracts (> 250,000 ILS) require human audit
        if missing_optional or (doc.deal_amount is not None and doc.deal_amount > 250000):
            routing_status = ProcessingStatus.NEEDS_REVIEW
        else:
            routing_status = ProcessingStatus.APPROVED

        return {
            "status": routing_status.value,
            "document": doc.model_dump(),
            "missing_fields": missing_optional,
            "errors": [],
        }

    except ValidationError as err:
        error_details = [
            {"field": " -> ".join(str(loc) for loc in e["loc"]), "message": e["msg"]}
            for e in err.errors()
        ]
        return {
            "status": ProcessingStatus.REJECTED.value,
            "document": None,
            "missing_fields": [],
            "errors": error_details,
        }


# ==========================================
# 3. Document File Ingestion (.txt, .pdf, .docx, images)
# ==========================================
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def load_document_text(file_path_or_text: str) -> str:
    """
    Universal ingestion function for text-based formats:
    - If input is a path to an existing .txt, .pdf, or .docx file, extracts the text.
    - If input is already raw text string, returns it directly.
    """
    if not os.path.isfile(file_path_or_text):
        return file_path_or_text

    ext = os.path.splitext(file_path_or_text)[1].lower()

    if ext == ".txt":
        with open(file_path_or_text, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == ".pdf":
        import pymupdf
        doc = pymupdf.open(file_path_or_text)
        text_parts: List[str] = [str(page.get_text()) for page in doc]
        return "\n".join(text_parts).strip()

    elif ext in (".docx", ".doc"):
        import docx
        doc = docx.Document(file_path_or_text)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs).strip()

    elif ext in IMAGE_EXTENSIONS:
        # Images will be uploaded directly to Gemini Multimodal Vision
        return ""

    else:
        raise ValueError(f"Unsupported file format '{ext}'. Expected .pdf, .docx, .txt, or image.")


# ==========================================
# 4. Live Gemini Multimodal API Caller
# ==========================================
def extract_with_gemini(
    document_source: str,
    api_key: Optional[str] = None,
    preferred_model: str = "gemini-3.5-flash-lite",
) -> Dict[str, Any]:
    """
    Invokes Google Gemini using structured outputs to guarantee
    schema alignment at generation time.
    Accepts:
    - File path (.pdf, .docx, .txt, or photos/scans in .png/.jpg)
    - Raw text string
    """
    key = api_key or os.getenv("GEMINI_API_KEY")
    if not key:
        raise ValueError("Missing GEMINI_API_KEY environment variable.")
    client = genai.Client(api_key=key)

    base_prompt = (
        "Extract vehicle document information from the provided document. "
        "For numeric amounts, provide raw numbers without currency symbols or commas."
    )

    contents: List[Any] = []

    # Check if input is an existing file
    if os.path.isfile(document_source):
        ext = os.path.splitext(document_source)[1].lower()

        # Case A: Visual photo or scan (PNG, JPG, WEBP)
        if ext in IMAGE_EXTENSIONS:
            uploaded_file = client.files.upload(file=document_source)
            contents = [uploaded_file, base_prompt]

        # Case B: PDF file (digital or scanned)
        elif ext == ".pdf":
            extracted_text = load_document_text(document_source)
            if len(extracted_text) >= 30:
                # Digital PDF with selectable text
                contents = [f"{base_prompt}\n\nDocument Text:\n{extracted_text}"]
            else:
                # Scanned image PDF -> Upload directly to Gemini Vision
                uploaded_file = client.files.upload(file=document_source)
                contents = [uploaded_file, base_prompt]

        # Case C: Word (.docx) or Text (.txt)
        else:
            extracted_text = load_document_text(document_source)
            contents = [f"{base_prompt}\n\nDocument Text:\n{extracted_text}"]
    else:
        # Case D: Raw text string
        contents = [f"{base_prompt}\n\nDocument Text:\n{document_source}"]

    candidate_models = [preferred_model, "gemini-3.1-flash-lite", "gemini-flash-lite-latest"]
    last_error = None

    for model_name in candidate_models:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=contents,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=ExtractedVehicleDocument,
                    temperature=0.0,
                ),
            )
            if isinstance(response.parsed, BaseModel):
                return response.parsed.model_dump()
            elif isinstance(response.parsed, dict):
                return response.parsed
            elif response.text:
                return json.loads(response.text)
        except Exception as e:
            last_error = e
            continue

    raise RuntimeError(f"All candidate Gemini models failed. Last error: {last_error}")


def process_document_source(
    document_source: str,
    api_key: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Unified end-to-end processing entry point:
    Accepts either a file path (.pdf, .docx, .txt) or a raw text string.
    1. Ingests file and extracts raw text.
    2. Invokes Gemini to extract structured fields.
    3. Validates and routes through business decision engine.
    """
    extracted_data = extract_with_gemini(document_source, api_key=api_key)
    return process_extracted_document(extracted_data)



# ==========================================
# 4. Self-Contained Test Suite & Demonstration
# ==========================================
if __name__ == "__main__":
    print("=" * 75)
    print("FREESBE DOCUMENT INTELLIGENCE PIPELINE - DEMONSTRATION")
    print("=" * 75)

    test_scenarios = [
        {
            "name": "Scenario 1: Clean, valid vehicle financing contract -> APPROVED",
            "data": {
                "document_id": "DOC-2026-001",
                "license_plate": "123-45-678",
                "customer_name": "Yossi Cohen",
                "customer_id": "012345674",  # Valid Israeli ID checksum
                "deal_amount": 142000.0,
                "contract_date": "2026-03-15",
                "category": "Financing Contract",
                "raw_text_snippet": "הסכם מימון רכב שנחתם במועד...",
            },
        },
        {
            "name": "Scenario 2: Missing optional date & high amount (>250k) -> NEEDS_REVIEW",
            "data": {
                "document_id": "DOC-2026-002",
                "license_plate": "9876543",
                "customer_name": "Rina Levi",
                "customer_id": None,
                "deal_amount": 340000.0,
                "contract_date": None,
                "category": None,
                "raw_text_snippet": "אישור הזמנת רכב יוקרה - לוח סילוקין",
            },
        },
        {
            "name": "Scenario 3: Corrupted OCR (Invalid plate & negative deal) -> REJECTED",
            "data": {
                "document_id": "DOC-2026-003",
                "license_plate": "INVALID-PLATE",
                "customer_name": "Avi Ron",
                "deal_amount": -5000.0,
                "contract_date": "2026-01-01",
                "raw_text_snippet": "רישיון רכב פרטי",
            },
        },
        {
            "name": "Scenario 4: Cross-Field Rule Violation (Damage Report missing estimated_damage) -> REJECTED",
            "data": {
                "document_id": "DOC-2026-004",
                "license_plate": "7788899",
                "customer_name": "Moshe Peretz",
                "category": "Damage Report",
                "estimated_damage": None,  # Violates cross-field requirement
                "raw_text_snippet": "דו\"ח שמאי רכב על תאונה בצד שמאל",
            },
        },
    ]

    for scenario in test_scenarios:
        print(f"\n[*] Executing: {scenario['name']}")
        result = process_extracted_document(scenario["data"])
        print(f"    Processing Status: {result['status']}")

        if result["status"] == ProcessingStatus.REJECTED.value:
            print(f"    Validation Errors: {result['errors']}")
        else:
            print(f"    Validated Payload: {result['document']}")
            print(f"    Missing Fields:    {result['missing_fields']}")
        print("-" * 75)

    # Scenario 5: Live API Call (if GEMINI_API_KEY is present)
    api_key = os.getenv("GEMINI_API_KEY")
    if api_key:
        print("\n[*] Executing: Scenario 5 (LIVE GEMINI EXTRACTION via gemini-3.5-flash-lite)")
        sample_doc = """
        טופס הסכם מימון והלוואה לרכב
        מספר אסמכתא: DOC-2026-LIVE-99
        מספר רישוי רכב: 23-456-78
        שם הלווה: דניאל קפלן
        תעודת זהות: 012345674
        סכום המימון המבוקש: 195000
        תאריך עריכת ההסכם: 2026-05-18
        """
        try:
            print("    Sending unstructured Hebrew text to Gemini...")
            extracted = extract_with_gemini(sample_doc, api_key=api_key)
            print("    Raw Structured Output from Gemini:")
            print(f"    -> {extracted}")
            
            pipeline_result = process_extracted_document(extracted)
            print(f"    Pipeline Routing Status: {pipeline_result['status']}")
            print(f"    Final Validated Payload: {pipeline_result['document']}")
        except Exception as e:
            print(f"    Live extraction skipped or failed: {e}")
        print("=" * 75)